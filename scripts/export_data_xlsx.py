#!/usr/bin/env python3
"""
将 DDRD / Conzzer / SegFuzz 8 个模块的实验数据导出到一个 xlsx 文件。
表格列: time_hours, count, module, tool

用法:
  python3 scripts/export_data_xlsx.py
  python3 scripts/export_data_xlsx.py -o my_output.xlsx
  python3 scripts/export_data_xlsx.py --metric varnames   # 只导出 varname pairs
  python3 scripts/export_data_xlsx.py --metric pairs      # 只导出 total pairs (默认)
  python3 scripts/export_data_xlsx.py --metric both       # 两种指标都导出
"""

import os
import re
import csv
import glob
import argparse
import sys
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("请先安装 openpyxl: pip install openpyxl")
    sys.exit(1)

# ── 目录配置 ──────────────────────────────────────────────
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
EXP_DIR = os.path.join(PROJECT_ROOT, "exp")

CONZZER_CSV_DIR = "/home/zzzccc/BASS/Conzzer/exp/kccwf_persistent_data"
SEGFUZZ_CSV_DIR = "/home/zzzccc/BASS/segfuzz/exp/segfuzz-kccwf-comparison/kccwf_persistent_data"
SEGFUZZ_TIME_SCALE = 2.0

TARGET_MODULES = ["xfs", "btrfs", "f2fs", "jfs", "floppy", "bt-stack", "ptmx", "dsp"]
PLOT_MAX_HOURS = 24.0

LOG_PATTERN = re.compile(
    r'(\d{4}/\d{2}/\d{2} \d{2}:\d{2}:\d{2})'
    r'.*?uaf pairs=(\d+)'
    r'.*?uaf varnames=(\d+)'
)
LOG_PATTERN_ALT = re.compile(
    r'(\d{4}/\d{2}/\d{2} \d{2}:\d{2}:\d{2})'
    r'.*?uaf varnames=(\d+)'
    r'.*?uaf pairs=(\d+)'
)


# ── 数据加载 ──────────────────────────────────────────────

def parse_ddrd_log(log_path):
    raw = []
    with open(log_path, 'r', encoding='utf-8', errors='ignore') as f:
        for line in f:
            m = LOG_PATTERN.search(line)
            if m:
                ts = datetime.strptime(m.group(1), '%Y/%m/%d %H:%M:%S')
                raw.append((ts, int(m.group(2)), int(m.group(3))))
                continue
            m = LOG_PATTERN_ALT.search(line)
            if m:
                ts = datetime.strptime(m.group(1), '%Y/%m/%d %H:%M:%S')
                raw.append((ts, int(m.group(3)), int(m.group(2))))
    if not raw:
        return [], [], []
    raw.sort(key=lambda x: x[0])
    t0 = raw[0][0]
    minute_data = {}
    for ts, pairs, varnames in raw:
        key = ts.replace(second=0, microsecond=0)
        minute_data[key] = (pairs, varnames)
    sorted_keys = sorted(minute_data.keys())
    hours = [(k - t0).total_seconds() / 3600 for k in sorted_keys]
    pairs_vals = [minute_data[k][0] for k in sorted_keys]
    varnames_vals = [minute_data[k][1] for k in sorted_keys]
    return hours, pairs_vals, varnames_vals


def find_latest_ddrd_log(module):
    log_dir = os.path.join(EXP_DIR, module, "logs")
    if not os.path.isdir(log_dir):
        return None
    candidates = [p for p in glob.glob(os.path.join(log_dir, "exp-fuzz*.log"))
                  if '-random-' not in os.path.basename(p)]
    if not candidates:
        return None
    candidates.sort(key=lambda p: os.path.getmtime(p), reverse=True)
    return candidates[0]


def load_conzzer_csv(module):
    csv_path = os.path.join(CONZZER_CSV_DIR, f"{module}.csv")
    if not os.path.isfile(csv_path):
        return [], [], []
    rows = []
    with open(csv_path, 'r') as f:
        reader = csv.DictReader(f)
        for row in reader:
            wall_clock = float(row['wall_clock'])
            elapsed_seconds = float(row['elapsed_min']) * 60.0
            rows.append((wall_clock, elapsed_seconds, int(row['capped_race_signals']), int(row['varname_pairs'])))
    if not rows:
        return [], [], []
    inferred_start = min(wc - es for wc, es, _, _ in rows)
    hours, signals, varnames = [], [], []
    for wc, _, sig, var in rows:
        hours.append((wc - inferred_start) / 3600.0)
        signals.append(sig)
        varnames.append(var)
    return hours, signals, varnames


def load_segfuzz_csv(module):
    csv_path = os.path.join(SEGFUZZ_CSV_DIR, f"{module}.csv")
    if not os.path.isfile(csv_path):
        return [], [], []
    hours, pairs, varnames = [], [], []
    with open(csv_path, 'r') as f:
        reader = csv.DictReader(f)
        for row in reader:
            h = float(row['elapsed_min']) / 60.0 / SEGFUZZ_TIME_SCALE
            hours.append(h)
            pairs.append(int(row['unique_race_pairs']))
            varnames.append(int(row['unique_varname_pairs']))
    return hours, pairs, varnames


def _make_monotonic(values):
    if not values:
        return values
    out = [values[0]]
    for v in values[1:]:
        out.append(max(out[-1], v))
    return out


def _first_effective_index(*value_lists):
    if not value_lists or not value_lists[0]:
        return None
    for i in range(len(value_lists[0])):
        if any(vl[i] > 0 for vl in value_lists):
            return i
    return None


def adjust_and_clip(hours, *value_lists,
                    warmup_hours=1.0, immediate_threshold=1.0, max_hours=PLOT_MAX_HOURS):
    """与 plot_4way.py 中 adjust_series_start + clip_series 等效的简化版本。"""
    if not hours:
        return (hours,) + value_lists

    first_idx = _first_effective_index(*value_lists)
    if first_idx is None:
        return (hours,) + value_lists

    first_h = hours[first_idx]
    if first_h <= immediate_threshold:
        shift = max(0.0, warmup_hours - first_h)
        hours = [h + shift for h in hours]
    else:
        hours = [h - first_h for h in hours]
        # 只保留 first_idx 之后的数据
        hours = hours[first_idx:]
        value_lists = tuple(vl[first_idx:] for vl in value_lists)

    # clip
    clipped_h = []
    clipped_v = [[] for _ in value_lists]
    for i, h in enumerate(hours):
        if h > max_hours:
            break
        clipped_h.append(h)
        for vi, vl in enumerate(value_lists):
            clipped_v[vi].append(vl[i])

    return (clipped_h,) + tuple(clipped_v)


# ── 样式辅助 ──────────────────────────────────────────────

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)

# 工具颜色（与图表配色一致）
TOOL_COLORS = {
    "DDRD":    "D6E4F0",  # 蓝色系浅色
    "Conzzer": "D5F0D6",  # 绿色系浅色
    "SegFuzz": "FFF3CD",  # 黄色系浅色
}

MODULE_ALT_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")


def style_header(ws, header_row):
    for col_idx, _ in enumerate(header_row, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def auto_column_width(ws):
    """根据表头名称设定列宽，避免遍历全表。"""
    for col in ws.iter_cols(min_row=1, max_row=1):
        cell = col[0]
        col_letter = get_column_letter(cell.column)
        # 根据表头名称估算宽度
        name = str(cell.value) if cell.value else ""
        width = max(len(name) + 4, 12)
        ws.column_dimensions[col_letter].width = width


def resample_to_minutes(hours, *value_lists):
    """将时间序列按整分钟降采样，每分钟取最后一个值（保持累计量语义）。
    返回 (hours_resampled, v1_resampled, v2_resampled, ...)，time 单位保持 hours。"""
    if not hours:
        return (hours,) + value_lists

    from math import floor
    # 按分钟分组，取每分钟最后一个值
    minute_data = {}  # minute_int -> (hour, v1, v2, ...)
    for i, h in enumerate(hours):
        minute = floor(h * 60)
        vals = tuple(vl[i] for vl in value_lists)
        minute_data[minute] = (h, vals)  # 后面的覆盖前面的（取最后值）

    sorted_minutes = sorted(minute_data.keys())
    out_hours = [minute_data[m][0] for m in sorted_minutes]
    # 用整分钟数换算成小时，便于表格阅读
    out_hours = [m / 60.0 for m in sorted_minutes]
    out_values = [[] for _ in value_lists]
    for m in sorted_minutes:
        vals = minute_data[m][1]
        for vi, v in enumerate(vals):
            out_values[vi].append(v)

    return (out_hours,) + tuple(out_values)


# ── 主流程 ────────────────────────────────────────────────

def collect_rows(metric):
    """
    metric: 'pairs' | 'varnames' | 'both'
    返回列表: [(time_hours, count, module, tool, metric_name), ...]
    每分钟一个数据点。
    """
    all_rows = []
    for module in TARGET_MODULES:
        print(f"  [{module}]", end=" ")

        # DDRD
        log_path = find_latest_ddrd_log(module)
        if log_path:
            h, pairs, varnames = parse_ddrd_log(log_path)
            h, pairs, varnames = _make_monotonic_series(h, pairs, varnames)
            result = adjust_and_clip(h, pairs, varnames)
            h, pairs, varnames = result[0], result[1], result[2]
            h, pairs, varnames = resample_to_minutes(h, pairs, varnames)
            if metric in ('pairs', 'both'):
                for th, tv in zip(h, pairs):
                    all_rows.append((round(th, 4), tv, module, "DDRD", "pairs"))
            if metric in ('varnames', 'both'):
                for th, tv in zip(h, varnames):
                    all_rows.append((round(th, 4), tv, module, "DDRD", "varnames"))
            print(f"DDRD({len(h)}pts)", end=" ")
        else:
            print("DDRD(miss)", end=" ")

        # Conzzer
        h, signals, varnames = load_conzzer_csv(module)
        if h:
            signals = _make_monotonic(signals)
            varnames = _make_monotonic(varnames)
            result = adjust_and_clip(h, signals, varnames)
            h, signals, varnames = result[0], result[1], result[2]
            h, signals, varnames = resample_to_minutes(h, signals, varnames)
            if metric in ('pairs', 'both'):
                for th, tv in zip(h, signals):
                    all_rows.append((round(th, 4), tv, module, "Conzzer", "pairs"))
            if metric in ('varnames', 'both'):
                for th, tv in zip(h, varnames):
                    all_rows.append((round(th, 4), tv, module, "Conzzer", "varnames"))
            print(f"Conzzer({len(h)}pts)", end=" ")
        else:
            print("Conzzer(miss)", end=" ")

        # SegFuzz
        h, pairs, varnames = load_segfuzz_csv(module)
        if h:
            pairs = _make_monotonic(pairs)
            varnames = _make_monotonic(varnames)
            result = adjust_and_clip(h, pairs, varnames)
            h, pairs, varnames = result[0], result[1], result[2]
            h, pairs, varnames = resample_to_minutes(h, pairs, varnames)
            if metric in ('pairs', 'both'):
                for th, tv in zip(h, pairs):
                    all_rows.append((round(th, 4), tv, module, "SegFuzz", "pairs"))
            if metric in ('varnames', 'both'):
                for th, tv in zip(h, varnames):
                    all_rows.append((round(th, 4), tv, module, "SegFuzz", "varnames"))
            print(f"SegFuzz({len(h)}pts)", end=" ")
        else:
            print("SegFuzz(miss)", end=" ")

        print()

    return all_rows


def _make_monotonic_series(hours, *value_lists):
    """对多个 value_list 同时做单调化，返回 (hours, v1, v2, ...)。"""
    return (hours,) + tuple(_make_monotonic(vl) for vl in value_lists)


def write_xlsx(all_rows, output_path, metric):
    wb = openpyxl.Workbook()

    if metric == 'both':
        header = ["time_hours", "count", "module", "tool", "metric"]
    else:
        header = ["time_hours", "count", "module", "tool"]

    # 预建填充对象（复用，不每行新建）
    _fills = {tool: PatternFill(start_color=c, end_color=c, fill_type="solid")
              for tool, c in TOOL_COLORS.items()}

    # ── Sheet 1: 每个模块一个 Sheet ───────────────────────
    first_sheet = True
    for module in TARGET_MODULES:
        if first_sheet:
            ws = wb.active
            ws.title = module
            first_sheet = False
        else:
            ws = wb.create_sheet(title=module)

        ws.append(header)
        style_header(ws, header)

        module_rows = [r for r in all_rows if r[2] == module]
        module_rows.sort(key=lambda r: (r[3], r[4], r[0]))

        cur_row = 2
        cur_tool = None
        for row_data in module_rows:
            th, count, module_name, tool, metric_name = row_data
            if metric == 'both':
                ws.cell(cur_row, 1, th)
                ws.cell(cur_row, 2, count)
                ws.cell(cur_row, 3, module_name)
                ws.cell(cur_row, 4, tool)
                ws.cell(cur_row, 5, metric_name)
            else:
                ws.cell(cur_row, 1, th)
                ws.cell(cur_row, 2, count)
                ws.cell(cur_row, 3, module_name)
                ws.cell(cur_row, 4, tool)
            cur_tool = tool
            cur_row += 1

        ws.freeze_panes = "A2"
        # 固定列宽
        col_widths = {"time_hours": 14, "count": 10, "module": 12,
                      "tool": 12, "metric": 12}
        for ci, col_name in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(ci)].width = col_widths.get(col_name, 12)

        print(f"  [{module}] {cur_row - 2} 行", flush=True)

    # ── Summary Sheet ──────────────────────────────────────
    ws_sum = wb.create_sheet(title="Summary")
    if metric == 'both':
        sum_header = ["module", "tool", "metric", "final_count", "max_count", "data_points"]
    else:
        sum_header = ["module", "tool", "final_count", "max_count", "data_points"]
    ws_sum.append(sum_header)
    style_header(ws_sum, sum_header)

    from collections import defaultdict
    groups = defaultdict(list)
    for th, count, module, tool, metric_name in all_rows:
        groups[(module, tool, metric_name)].append((th, count))

    cur_row = 2
    for module in TARGET_MODULES:
        for tool in ["DDRD", "Conzzer", "SegFuzz"]:
            metrics_list = ["pairs", "varnames"] if metric == 'both' else [metric]
            for met in metrics_list:
                key = (module, tool, met)
                if key not in groups:
                    continue
                pts = sorted(groups[key], key=lambda x: x[0])
                counts = [c for _, c in pts]
                final_count = counts[-1] if counts else 0
                max_count = max(counts) if counts else 0
                n_pts = len(counts)
                fill = _fills.get(tool)
                if metric == 'both':
                    values = [module, tool, met, final_count, max_count, n_pts]
                else:
                    values = [module, tool, final_count, max_count, n_pts]
                for ci, v in enumerate(values, 1):
                    cell = ws_sum.cell(cur_row, ci, v)
                    if fill:
                        cell.fill = fill
                cur_row += 1

    ws_sum.freeze_panes = "A2"
    for ci, col_name in enumerate(sum_header, 1):
        ws_sum.column_dimensions[get_column_letter(ci)].width = 16

    wb.save(output_path)
    print(f"\n已保存: {output_path}")
    print(f"共 {len(all_rows)} 行数据，{len(TARGET_MODULES)} 个模块 Sheet + Summary")


def _darken(hex_color, factor=0.92):
    """将颜色略微加深用于交替行。"""
    r = int(hex_color[0:2], 16)
    g = int(hex_color[2:4], 16)
    b = int(hex_color[4:6], 16)
    r = int(r * factor)
    g = int(g * factor)
    b = int(b * factor)
    return f"{r:02X}{g:02X}{b:02X}"


def main():
    parser = argparse.ArgumentParser(description="导出实验数据到 xlsx")
    parser.add_argument("-o", "--output", default=None,
                        help="输出文件路径 (默认: exp/4way_compare/data_export.xlsx)")
    parser.add_argument("--metric", choices=["pairs", "varnames", "both"], default="pairs",
                        help="导出哪种指标 (默认: pairs)")
    args = parser.parse_args()

    if args.output is None:
        out_dir = os.path.join(PROJECT_ROOT, "exp", "4way_compare")
        os.makedirs(out_dir, exist_ok=True)
        args.output = os.path.join(out_dir, "data_export.xlsx")

    print(f"指标: {args.metric}")
    print("加载数据...")
    all_rows = collect_rows(args.metric)
    print(f"\n写入 xlsx ...")
    write_xlsx(all_rows, args.output, args.metric)


if __name__ == "__main__":
    main()
